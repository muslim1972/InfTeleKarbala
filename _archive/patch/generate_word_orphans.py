# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page Setup - Landscape for wide table
    for section in doc.sections:
        section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Header
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_head = p_head.add_run("جمهورية العراق - وزارة الاتصالات\nالشركة العامة للاتصالات والمعلوماتية\nمديرية اتصالات ومعلوماتية كربلاء المقدسة / شعبة الموارد البشرية")
    r_head.font.name = "Arial"
    r_head.font.size = Pt(13)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(0, 51, 102)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("استمارة حصر وتحديث التشكيل الإداري وموقع الدوام للمنتسبين غير المحددين")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(180, 0, 0)

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_note = p_note.add_run("ملاحظة للإدارة: يرجى التفضل بتحديد (القسم) و(الشعبة/المجمع) و(موقع الدوام المعتمد للبصمة) لكل منتسب في الحقول المخصصة أدناه، ليتم إدخالها وتحديث ارتباطهم الهرمي في النظام بشكل نهائي.")
    r_note.font.name = "Arial"
    r_note.font.size = Pt(10)
    r_note.font.italic = True
    r_note.font.color.rgb = RGBColor(80, 80, 80)

    employees = [
        ("1", "143264419", "ابراهيم عبد الامير عراك", "حرفي", "الابتدائية"),
        ("2", "143264323", "امير علاوي نعمة", "حرفي", "الابتدائية"),
        ("3", "c123456", "جلنار زهير", "غير محدد", "غير محدد"),
        ("4", "143264464", "حسن عبد علي سالم", "حرفي", "المتوسطة"),
        ("5", "143264310", "حسين علي سعيد", "حرفي", "الاعدادية"),
        ("6", "143264381", "حيدر حميد حسين", "حرفي", "الابتدائية"),
        ("7", "143264253", "سجاد عبد اليمة فليح", "حرفي", "الابتدائية"),
        ("8", "143264282", "علي حسين عبد عون", "حرفي", "المتوسطة"),
        ("9", "150975559", "علي صاحب حسين", "محاسب", "ماجستير"),
        ("10", "143264448", "علي محمد عبد الرضا", "حرفي", "الابتدائية"),
        ("11", "143264435", "كرار حيدر عبد الامام", "حرفي", "الابتدائية"),
        ("12", "143264422", "محمد احمد عبد الحسين", "حرفي", "المتوسطة"),
        ("13", "143264505", "محمد صادق جواد", "حرفي", "الابتدائية"),
        ("14", "143264295", "محمد عباس محسن", "حرفي", "الابتدائية"),
        ("15", "143264349", "محمود نعمة شنين", "حرفي", "الابتدائية"),
        ("16", "143264240", "مصطفى حيدر ياسين", "حرفي", "الابتدائية"),
        ("17", "143264378", "مصطفى طالب ابراهيم", "حرفي", "الابتدائية"),
        ("18", "143264365", "نهضة عبد الله كاظم", "حرفي", "الابتدائية"),
        ("19", "143264394", "وحيد فاهم وحيد", "حرفي", "الاعدادية"),
        ("20", "143264406", "وسيم عباس صبري", "حرفي", "الابتدائية"),
    ]

    headers = ["ت", "الرقم الوظيفي", "الاسم الكامل للمنتسب", "العنوان", "الشهادة", "القسم المطلوب", "الشعبة / المجمع", "موقع البصمة / الدوام", "توقيع / ملاحظات الموارد"]
    
    table = doc.add_table(rows=len(employees) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(0.4), Inches(1.1), Inches(2.2), Inches(0.9), Inches(1.0), Inches(1.5), Inches(1.6), Inches(1.6), Inches(1.5)]

    # Style Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1F4E78") # Deep Blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Style Data Rows
    for row_idx, emp in enumerate(employees, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F2F5F8" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx in range(len(headers)):
            cell = row_cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            if col_idx < len(emp):
                cell.text = emp[col_idx]
            else:
                cell.text = "" # Blank for HR to write in

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 2 else WD_ALIGN_PARAGRAPH.RIGHT
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                if col_idx == 0:
                    r.font.bold = True
                elif col_idx == 1:
                    r.font.name = "Consolas"
                    r.font.bold = True

    # Set column widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Footer Signatures
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(24)
    table_sig = doc.add_table(rows=1, cols=3)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_widths = [Inches(3.5), Inches(3.5), Inches(3.5)]
    for row in table_sig.rows:
        for i, w in enumerate(sig_widths):
            row.cells[i].width = w

    c1 = table_sig.rows[0].cells[0].paragraphs[0]
    c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = c1.add_run("مسؤول وحدة الإضابير\n\n__________________")
    r1.font.bold = True
    r1.font.name = "Arial"

    c2 = table_sig.rows[0].cells[1].paragraphs[0]
    c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = c2.add_run("مسؤول شعبة الموارد البشرية\n\n__________________")
    r2.font.bold = True
    r2.font.name = "Arial"

    c3 = table_sig.rows[0].cells[2].paragraphs[0]
    c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = c3.add_run("مصادقة السيد المعاون / المدير\n\n__________________")
    r3.font.bold = True
    r3.font.name = "Arial"

    out_file = r"F:\صور للتجربة\نظام المديرية\نظام الادارة\تطبيق نحن\استمارة_تحديد_التشكيلات_الإدارية_للمنتسبين.docx"
    doc.save(out_file)
    print("Successfully generated Word DOCX at:", out_file)

if __name__ == "__main__":
    create_document()
