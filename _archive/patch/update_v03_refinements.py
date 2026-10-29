# -*- coding: utf-8 -*-
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy

HIGHLIGHT_COLOR = RGBColor(0, 82, 204) # Royal Blue for v0.3 updates

file_path = r"F:\صور للتجربة\نظام المديرية\نظام الادارة\تطبيق نحن\البحث النهائي\بحث إطار عمل B.R.I.C.K ومحرك الإنفاذ الحتمي v0.3.docx"
doc = docx.Document(file_path)

# =========================================================================
# 1. FIX THE FORMULA FOR INVARIANT 6 (Question 1 & Image 1)
# =========================================================================
# Find the raw formula paragraph and paragraph 118
p_formula = None
p_inv6 = None
for i, p in enumerate(doc.paragraphs):
    if "مسلمة التأصيل الهيكلي القطعي" in p.text:
        p_inv6 = p
    if "Columns(f)" in p.text and "DryRun(f)" in p.text:
        p_formula = p

if p_inv6 and p_formula:
    print("Found Invariant 6 text and raw formula paragraph.")
    # Clone Table 5 (which is the formula box for Invariant 4)
    t5 = doc.tables[5]
    new_tbl_element = copy.deepcopy(t5._tbl)
    
    # Insert new table right after p_inv6
    p_inv6._p.addnext(new_tbl_element)
    
    # Wrap new table in docx Table object
    new_tbl = docx.table.Table(new_tbl_element, doc)
    
    # Clear cell content and set proper Consolas runs
    cell = new_tbl.cell(0, 0)
    cell.text = ""
    p_cell = cell.paragraphs[0]
    p_cell.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    
    # Formula components with Consolas font
    formula_parts = [
        ("∀ f ∈ ", True),
        ("S_rpc", True),
        (",  Columns(f) ⊆ Catalog(", True),
        ("S_tables", True),
        (")  ∧  DryRun(f) ≠ RuntimeSchemaError", True)
    ]
    
    for text_part, is_bold in formula_parts:
        r = p_cell.add_run(text_part)
        r.font.name = "Consolas"
        r.font.bold = is_bold
        r.font.color.rgb = HIGHLIGHT_COLOR
        # Ensure Word uses Consolas for complex scripts as well
        rPr = r._r.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Consolas')
        rFonts.set(qn('w:hAnsi'), 'Consolas')
        rFonts.set(qn('w:cs'), 'Consolas')
        rPr.append(rFonts)
    
    # Remove raw unboxed paragraph
    p_formula._p.getparent().remove(p_formula._p)
    print("Replaced raw formula paragraph with styled 1x1 formula table using Consolas.")

# =========================================================================
# 2. UPDATE TABLE (Question 2 & Image 2)
# =========================================================================
# Find table with 'المقياس الهندسي'
t_scale = None
for tbl in doc.tables:
    if len(tbl.rows) > 0 and "المقياس الهندسي" in tbl.cell(0, 0).text:
        t_scale = tbl
        break

if t_scale:
    cell_fingerprints = t_scale.cell(2, 0) # wait, let's check which cell has 'مع 386'
    for r in t_scale.rows:
        for c in r.cells:
            if "مسجلاً بالبصمة" in c.text:
                c.text = ""
                p_fp = c.paragraphs[0]
                r_fp1 = p_fp.add_run("تغطية بنسبة 100% (مع ")
                r_fp2 = p_fp.add_run("400 ملفاً مسجلاً بالبصمة التشفيرية")
                r_fp2.font.color.rgb = HIGHLIGHT_COLOR
                r_fp2.font.bold = True
                r_fp3 = p_fp.add_run(")")
                print("Updated Table Scale with 400 fingerprints.")
                break

# =========================================================================
# 3. UPDATE SECTION 6.2 PERFORMANCE BENCHMARK NUMBERS (Question 3 & Image 3)
# =========================================================================
# Update paragraphs 200 and 201 (or search by content)
for p in doc.paragraphs:
    if "زمن التحقق الهيكلي المحلي ومستودع الظل:" in p.text:
        p.text = ""
        r_p1 = p.add_run("▪ زمن التحقق الهيكلي المحلي ومستودع الظل: ")
        r_p2 = p.add_run("بلغ متوسط زمن تنفيذ الفحص المحلي والتحقق التشفيري 1.16 ثانية لفحص ")
        r_p3 = p.add_run("318 ملفاً مصدرياً وتدقيق 400 بصمة تجزئة تشفيرية")
        r_p3.font.color.rgb = HIGHLIGHT_COLOR
        r_p3.font.bold = True
        r_p4 = p.add_run("، فيما بلغ زمن استرجاع الملف التالف من مستودع الظل أقل من 5 أجزاء من الألف من الثانية (0.005s).")
        print("Updated Section 6.2 Local Overhead paragraph.")
        
    if "زمن الفحص المباشر لقاعدة البيانات البعيدة:" in p.text:
        p.text = ""
        r_d1 = p.add_run("▪ زمن الفحص المباشر لقاعدة البيانات البعيدة: ")
        r_d2 = p.add_run("استغرق فحص الخادم الحي عبر بروتوكول آمن متوسط 3.42 ثانية، شمل استعلام مخططات pg_proc وinformation_schema ومقارنة ")
        r_d3 = p.add_run("228 كائناً")
        r_d3.font.color.rgb = HIGHLIGHT_COLOR
        r_d3.font.bold = True
        r_d4 = p.add_run(" موزعة على المخططين public وitpc (128 دالة RPC مخزنة، 91 جدولاً، 9 حاويات تخزين).")
        print("Updated Section 6.2 DB Introspection Overhead paragraph.")

doc.save(file_path)
print("v0.3.docx updated and saved successfully.")
